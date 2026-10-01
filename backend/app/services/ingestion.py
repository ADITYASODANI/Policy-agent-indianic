"""Policy ingestion: load PDF/DOCX/TXT -> split by section -> embed -> ChromaDB."""

import logging
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from app.config import get_settings
from app.services.vectorstore import reset_vectorstore

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

# Top-level headings look like "4. EMPLOYEE HEALTH, SAFETY & WORKPLACE SECURITY".
_CAPS_HEADING = re.compile(r"^\s*(\d{1,2})\.\s+([A-Z][A-Z0-9 &,/()'’\-–—]{2,120})\s*$")
# Fallback for documents whose headings are not upper-case.
_ANY_HEADING = re.compile(r"^\s*(\d{1,2})\.\s+([A-Z][^.\n]{2,100})\s*$")

_SMALL_WORDS = {"a", "an", "and", "at", "by", "for", "from", "in", "of", "on", "or", "the", "to"}

OVERVIEW_SECTION = "0"
OVERVIEW_TITLE = "Policy Information"


class IngestionError(Exception):
    pass


@dataclass
class Section:
    number: str
    title: str
    text: str


def load_text(path: Path) -> str:
    """Extract raw text from a PDF, DOCX or TXT file."""
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise IngestionError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")
    if not path.is_file():
        raise IngestionError(f"Policy file not found: {path}")

    try:
        if ext == ".pdf":
            reader = PdfReader(str(path))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == ".docx":
            doc = DocxDocument(str(path))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        else:
            text = path.read_text(encoding="utf-8")
    except Exception as exc:  # pypdf/python-docx raise a variety of errors
        raise IngestionError(f"Could not read policy file: {exc}") from exc

    if not text.strip():
        raise IngestionError("No text could be extracted from the policy file.")
    return text


def clean_text(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    text = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def title_case(heading: str) -> str:
    """'REMOTE WORK & WORK-FROM-HOME (WFH) RESPONSIBILITIES' -> 'Remote Work & Work-From-Home (WFH) Responsibilities'."""
    words = []
    for i, word in enumerate(heading.split()):
        if word.startswith("(") and word.endswith(")"):
            words.append(word.upper())  # acronyms such as (POSH), (WFH)
        elif i > 0 and word.lower() in _SMALL_WORDS:
            words.append(word.lower())
        else:
            words.append("-".join(part.capitalize() for part in word.split("-")))
    return " ".join(words)


def extract_metadata(text: str) -> dict[str, str]:
    settings = get_settings()
    meta = {
        "policy_name": settings.policy_name,
        "version": settings.policy_version,
        "effective_date": settings.policy_effective_date,
    }
    if m := re.search(r"Effective Date\s*:\s*(.+)", text, re.IGNORECASE):
        meta["effective_date"] = m.group(1).strip()
    if m := re.search(r"\bVersion\s*(\d+(?:\.\d+)?)", text, re.IGNORECASE):
        meta["version"] = f"Version {m.group(1)}"
    return meta


def split_sections(text: str) -> list[Section]:
    lines = text.split("\n")
    pattern = _CAPS_HEADING
    if not any(_CAPS_HEADING.match(line) for line in lines):
        pattern = _ANY_HEADING

    sections: list[Section] = []
    current = Section(OVERVIEW_SECTION, OVERVIEW_TITLE, "")
    buffer: list[str] = []

    def flush() -> None:
        body = "\n".join(buffer).strip()
        if body or current.number != OVERVIEW_SECTION:
            sections.append(Section(current.number, current.title, body))

    for line in lines:
        match = pattern.match(line)
        if match:
            flush()
            title = match.group(2).strip()
            if title.isupper():
                title = title_case(title)
            current = Section(match.group(1), title, "")
            buffer = []
        else:
            buffer.append(line)
    flush()

    if len([s for s in sections if s.number != OVERVIEW_SECTION]) == 0:
        raise IngestionError("No numbered policy sections were detected in the document.")
    return sections


def build_chunks(sections: list[Section], base_meta: dict[str, str], source: str) -> list[Document]:
    settings = get_settings()
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    docs: list[Document] = []
    for section in sections:
        if not section.text:
            continue
        heading = (
            section.title
            if section.number == OVERVIEW_SECTION
            else f"Section {section.number} — {section.title}"
        )
        for i, piece in enumerate(splitter.split_text(section.text)):
            # The heading is prepended so every chunk carries its section context
            # for both embedding similarity and the LLM prompt.
            docs.append(
                Document(
                    page_content=f"{heading}\n{piece}",
                    metadata={
                        **base_meta,
                        "section_number": section.number,
                        "section_title": section.title,
                        "chunk_index": i,
                        "source": source,
                    },
                )
            )
    return docs


def ingest_policy(path: Path | None = None) -> dict:
    """Run the full ingestion pipeline and replace the existing policy vectors."""
    path = Path(path or get_settings().policy_path)
    text = clean_text(load_text(path))
    meta = extract_metadata(text)
    sections = split_sections(text)
    chunks = build_chunks(sections, meta, path.name)

    store = reset_vectorstore()
    ids = [
        f"s{d.metadata['section_number']}-c{d.metadata['chunk_index']}" for d in chunks
    ]
    store.add_documents(chunks, ids=ids)
    logger.info("Ingested %d chunks from %s", len(chunks), path.name)

    return {
        "source": path.name,
        "chunks": len(chunks),
        "sections": [
            {"section_number": s.number, "section_title": s.title}
            for s in sections
            if s.number != OVERVIEW_SECTION
        ],
    }
